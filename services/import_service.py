"""
批量导入服务模块
支持：Markdown (.md)、文本 (.txt)、Excel (.xlsx) 导入
"""
import re
from pathlib import Path
from typing import List, Dict, Optional, Tuple
from dataclasses import dataclass

from models.account import Account
from models.url_item import URLItem


@dataclass
class ImportItem:
    """导入项（通用）"""
    app_name: str = ""
    username: str = ""
    password: str = ""
    url: str = ""
    category: str = ""
    tags: str = ""
    remark: str = ""
    valid: bool = True
    error_msg: str = ""
    
    def to_account(self) -> Account:
        import json
        tags_list = []
        if self.tags:
            # 支持逗号分隔或JSON数组格式
            try:
                parsed = json.loads(self.tags)
                if isinstance(parsed, list):
                    tags_list = parsed
            except (json.JSONDecodeError, TypeError):
                tags_list = [t.strip() for t in self.tags.split(',') if t.strip()]
        return Account(
            app_name=self.app_name,
            username=self.username,
            password=self.password,
            url=self.url,
            category=self.category or "其他",
            tags=json.dumps(tags_list, ensure_ascii=False) if tags_list else '[]',
            remark=self.remark
        )


# 字段名映射（支持中英文）
FIELD_MAPPING = {
    'app_name': ['应用名', 'app_name', 'App', '应用', '名称', 'Name', 'app', 'title', '标题'],
    'username': ['账号', 'username', '用户名', 'Account', 'User', '登录名', 'account', 'user', 'login'],
    'password': ['密码', 'password', '口令', 'Password', 'Pass', 'Pwd', 'pass', 'pwd'],
    'url': ['网址', 'url', 'URL', '网站', '链接', 'Link', 'website', 'address', '地址'],
    'category': ['分类', 'category', '类别', 'Category', '类型', 'Type', 'cat', 'group', '分组'],
    'remark': ['备注', 'remark', '说明', 'Remark', '描述', 'Description', 'desc', 'note', '注释'],
    'tags': ['标签', 'tags', 'tag', 'Tag', 'Tags', '标记', 'label', 'Label']
}


def _normalize_field_name(name: str) -> Optional[str]:
    """将各种字段名映射到标准字段名"""
    name_lower = name.strip().lower()
    for standard, aliases in FIELD_MAPPING.items():
        if name_lower in [a.lower() for a in aliases]:
            return standard
    return None


def _extract_md_value(text: str, prefixes: List[str]) -> str:
    """从 Markdown 行中提取值"""
    for prefix in prefixes:
        pattern = rf'^[-*]\s*{re.escape(prefix)}[:：]\s*(.*)$'
        match = re.search(pattern, text.strip(), re.IGNORECASE | re.MULTILINE)
        if match:
            return match.group(1).strip()
    return ""


def _strip_markdown_formatting(text: str) -> str:
    """去除 Markdown 格式标记（如 **bold**, *italic*, [link](url) 等）"""
    if not text:
        return text
    
    # 去除粗体 **text** 和 __text__
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    
    # 去除斜体 *text* 和 _text_
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    
    # 去除删除线 ~~text~~
    text = re.sub(r'~~(.+?)~~', r'\1', text)
    
    # 去除行内代码 `code`
    text = re.sub(r'`(.+?)`', r'\1', text)
    
    # 提取链接文本 [text](url) -> text
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    
    # 去除 HTML 标签 <tag>text</tag>
    text = re.sub(r'<[^>]+>', '', text)
    
    return text.strip()


class MarkdownParser:
    """Markdown 文件解析器（支持表格格式和段落格式）"""
    
    @staticmethod
    def parse(file_path: str) -> List[ImportItem]:
        content = Path(file_path).read_text(encoding='utf-8')
        
        # 检测是否为表格格式（包含 | 分隔的行）
        lines = content.split('\n')
        table_lines = [line for line in lines if line.strip().startswith('|')]
        
        if len(table_lines) >= 2:
            # 检测到表格格式
            return MarkdownParser._parse_table(content)
        else:
            # 使用段落格式解析
            return MarkdownParser._parse_sections(content)
    
    @staticmethod
    def _parse_table(content: str) -> List[ImportItem]:
        """解析 Markdown 表格格式"""
        items = []
        lines = content.split('\n')
        
        # 找到表格行
        table_lines = []
        for line in lines:
            line = line.strip()
            if line.startswith('|') and line.endswith('|'):
                table_lines.append(line)
        
        if len(table_lines) < 2:
            return items
        
        # 解析表头
        header_line = table_lines[0]
        headers = [cell.strip() for cell in header_line.strip('|').split('|')]
        
        # 映射列索引
        column_map = {}
        for idx, header in enumerate(headers):
            field = _normalize_field_name(header)
            if field:
                column_map[field] = idx
        
        # 检查必要列
        has_app_name = 'app_name' in column_map
        has_username = 'username' in column_map
        has_password = 'password' in column_map
        
        # 解析数据行（跳过分隔行 |---|---|）
        for line in table_lines[2:]:
            # 跳过分隔行
            if set(line.strip()).issubset({'|', '-', ':', ' '}):
                continue
            
            cells = [cell.strip() for cell in line.strip('|').split('|')]
            
            item = ImportItem()
            
            if has_app_name and column_map['app_name'] < len(cells):
                item.app_name = _strip_markdown_formatting(cells[column_map['app_name']])
            if has_username and column_map['username'] < len(cells):
                item.username = _strip_markdown_formatting(cells[column_map['username']])
            if has_password and column_map['password'] < len(cells):
                item.password = _strip_markdown_formatting(cells[column_map['password']])
            if 'url' in column_map and column_map['url'] < len(cells):
                item.url = _strip_markdown_formatting(cells[column_map['url']])
            if 'category' in column_map and column_map['category'] < len(cells):
                item.category = _strip_markdown_formatting(cells[column_map['category']])
            if 'remark' in column_map and column_map['remark'] < len(cells):
                item.remark = _strip_markdown_formatting(cells[column_map['remark']])
            if 'tags' in column_map and column_map['tags'] < len(cells):
                item.tags = _strip_markdown_formatting(cells[column_map['tags']])
            
            # 验证
            if not item.app_name or not item.username or not item.password:
                item.valid = False
                missing = []
                if not item.app_name:
                    missing.append('应用名')
                if not item.username:
                    missing.append('账号')
                if not item.password:
                    missing.append('密码')
                item.error_msg = f"缺少字段：{', '.join(missing)}"
            
            # 只添加非空行
            if item.app_name or item.username or item.password:
                items.append(item)
        
        return items
    
    @staticmethod
    def _parse_sections(content: str) -> List[ImportItem]:
        """解析段落式 Markdown（按 ## 标题分割）"""
        items = []
        
        # 按 ## 标题分割段落
        sections = re.split(r'\n##\s+', content)
        
        for section in sections:
            section = section.strip()
            if not section:
                continue
            
            lines = section.split('\n')
            app_name = lines[0].strip().lstrip('#').strip()
            if not app_name:
                continue
            
            body = '\n'.join(lines[1:])
            
            item = ImportItem()
            item.app_name = _strip_markdown_formatting(app_name)
            item.username = _strip_markdown_formatting(_extract_md_value(body, FIELD_MAPPING['username']))
            item.password = _strip_markdown_formatting(_extract_md_value(body, FIELD_MAPPING['password']))
            item.url = _strip_markdown_formatting(_extract_md_value(body, FIELD_MAPPING['url']))
            item.category = _strip_markdown_formatting(_extract_md_value(body, FIELD_MAPPING['category']))
            item.remark = _strip_markdown_formatting(_extract_md_value(body, FIELD_MAPPING['remark']))
            item.tags = _strip_markdown_formatting(_extract_md_value(body, FIELD_MAPPING['tags']))
            
            # 验证
            if not item.app_name or not item.username or not item.password:
                item.valid = False
                missing = []
                if not item.app_name:
                    missing.append('应用名')
                if not item.username:
                    missing.append('账号')
                if not item.password:
                    missing.append('密码')
                item.error_msg = f"缺少字段：{', '.join(missing)}"
            
            items.append(item)
        
        return items


class TextParser:
    """文本文件解析器（支持多种分隔符格式）"""
    
    @staticmethod
    def parse(file_path: str) -> List[ImportItem]:
        content = Path(file_path).read_text(encoding='utf-8')
        lines = content.split('\n')
        
        items = []
        current_item: Dict[str, str] = {}
        
        for line in lines:
            line = line.strip()
            if not line:
                # 空行表示一个记录结束
                if current_item:
                    item = TextParser._build_item(current_item)
                    items.append(item)
                    current_item = {}
                continue
            
            # 尝试等号分隔
            if '=' in line and not line.startswith('http'):
                parts = line.split('=', 1)
                if len(parts) == 2:
                    key, value = parts[0].strip(), parts[1].strip()
                    field = _normalize_field_name(key)
                    if field:
                        current_item[field] = value
                        continue
            
            # 尝试冒号分隔
            if ':' in line or '：' in line:
                parts = re.split(r'[:：]', line, 1)
                if len(parts) == 2:
                    key, value = parts[0].strip(), parts[1].strip()
                    field = _normalize_field_name(key)
                    if field:
                        current_item[field] = value
                        continue
            
            # 如果当前没有应用名，第一行作为应用名
            if 'app_name' not in current_item:
                current_item['app_name'] = line
            elif 'username' not in current_item:
                current_item['username'] = line
            elif 'password' not in current_item:
                current_item['password'] = line
            elif 'url' not in current_item:
                if line.startswith('http'):
                    current_item['url'] = line
                elif 'category' not in current_item:
                    current_item['category'] = line
                else:
                    current_item['remark'] = line
            elif 'category' not in current_item:
                current_item['category'] = line
            else:
                current_item['remark'] = line
        
        # 处理最后一个记录
        if current_item:
            item = TextParser._build_item(current_item)
            items.append(item)
        
        return items
    
    @staticmethod
    def _build_item(data: Dict[str, str]) -> ImportItem:
        item = ImportItem()
        item.app_name = _strip_markdown_formatting(data.get('app_name', ''))
        item.username = _strip_markdown_formatting(data.get('username', ''))
        item.password = _strip_markdown_formatting(data.get('password', ''))
        item.url = _strip_markdown_formatting(data.get('url', ''))
        item.category = _strip_markdown_formatting(data.get('category', ''))
        item.remark = _strip_markdown_formatting(data.get('remark', ''))
        item.tags = _strip_markdown_formatting(data.get('tags', ''))
        
        if not item.app_name or not item.username or not item.password:
            item.valid = False
            missing = []
            if not item.app_name:
                missing.append('应用名')
            if not item.username:
                missing.append('账号')
            if not item.password:
                missing.append('密码')
            item.error_msg = f"缺少字段：{', '.join(missing)}"
        
        return item


class ExcelParser:
    """Excel 文件解析器"""
    
    @staticmethod
    def parse(file_path: str) -> List[ImportItem]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("请安装 openpyxl：pip install openpyxl")
        
        wb = load_workbook(file_path)
        ws = wb.active
        
        # 读取表头
        headers = []
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        for cell in header_row:
            headers.append(str(cell) if cell else "")
        
        # 映射列索引
        column_map = {}
        for idx, header in enumerate(headers):
            field = _normalize_field_name(header)
            if field:
                column_map[field] = idx
        
        # 检查必要列
        has_app_name = 'app_name' in column_map
        has_username = 'username' in column_map
        has_password = 'password' in column_map
        
        items = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == '' for v in row):
                continue
            
            item = ImportItem()
            
            if has_app_name:
                item.app_name = str(row[column_map['app_name']]) if row[column_map['app_name']] is not None else ""
            if has_username:
                item.username = str(row[column_map['username']]) if row[column_map['username']] is not None else ""
            if has_password:
                item.password = str(row[column_map['password']]) if row[column_map['password']] is not None else ""
            if 'url' in column_map:
                item.url = str(row[column_map['url']]) if row[column_map['url']] is not None else ""
            if 'category' in column_map:
                item.category = str(row[column_map['category']]) if row[column_map['category']] is not None else ""
            if 'remark' in column_map:
                item.remark = str(row[column_map['remark']]) if row[column_map['remark']] is not None else ""
            if 'tags' in column_map:
                item.tags = str(row[column_map['tags']]) if row[column_map['tags']] is not None else ""
            
            # 验证
            if not item.app_name or not item.username or not item.password:
                item.valid = False
                missing = []
                if not item.app_name:
                    missing.append('应用名')
                if not item.username:
                    missing.append('账号')
                if not item.password:
                    missing.append('密码')
                item.error_msg = f"缺少字段：{', '.join(missing)}"
            
            items.append(item)
        
        return items


class URLParser:
    """网址批量导入解析器"""
    
    @staticmethod
    def parse_text(file_path: str) -> List[URLItem]:
        """从文本文件导入网址（每行一个 URL）"""
        content = Path(file_path).read_text(encoding='utf-8')
        lines = content.split('\n')
        
        items = []
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # 支持格式：URL 或 标题|URL
            if '|' in line:
                parts = line.split('|', 1)
                title = parts[0].strip()
                url = parts[1].strip()
            else:
                url = line
                title = ""
            
            # 简单验证 URL
            if not url.startswith(('http://', 'https://')):
                url = 'https://' + url
            
            items.append(URLItem(title=title, url=url))
        
        return items
    
    @staticmethod
    def parse_excel(file_path: str) -> List[URLItem]:
        """从 Excel 导入网址"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("请安装 openpyxl：pip install openpyxl")
        
        wb = load_workbook(file_path)
        ws = wb.active
        
        # 读取表头
        headers = []
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        for cell in header_row:
            headers.append(str(cell) if cell else "")
        
        # 映射列
        column_map = {}
        for idx, header in enumerate(headers):
            field = _normalize_field_name(header)
            if field:
                column_map[field] = idx
        
        items = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            
            url = ""
            title = ""
            category = ""
            
            if 'url' in column_map and row[column_map['url']]:
                url = str(row[column_map['url']])
            if 'app_name' in column_map and row[column_map['app_name']]:
                title = str(row[column_map['app_name']])
            if 'category' in column_map and row[column_map['category']]:
                category = str(row[column_map['category']])
            
            if url:
                if not url.startswith(('http://', 'https://')):
                    url = 'https://' + url
                items.append(URLItem(title=title, url=url, category=category))
        
        return items
    
    @staticmethod
    def parse_html_bookmarks(file_path: str) -> List[URLItem]:
        """
        解析浏览器收藏夹 HTML 文件（NETSCAPE-Bookmark-file-1 格式）
        支持 Chrome/Edge/Firefox 等浏览器导出的收藏夹
        """
        from html.parser import HTMLParser
        
        class BookmarkParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.items = []
                self.current_category = ""
                self.in_h3 = False
                self.h3_text = ""
                self.current_a = None
                self.current_a_text = ""
                
            def handle_starttag(self, tag, attrs):
                attrs_dict = dict(attrs)
                
                if tag == 'h3':
                    self.in_h3 = True
                    self.h3_text = ""
                elif tag == 'a':
                    # 提取链接信息
                    href = attrs_dict.get('href', '')
                    if href and href.startswith(('http://', 'https://')):
                        self.current_a = {
                            'url': href,
                            'title': attrs_dict.get('title', ''),
                            'category': self.current_category
                        }
                        self.current_a_text = ""
                        
            def handle_endtag(self, tag):
                if tag == 'h3':
                    self.in_h3 = False
                    # 使用 H3 作为分类名
                    if self.h3_text:
                        self.current_category = self.h3_text.strip()
                elif tag == 'a' and self.current_a:
                    # 如果没有 title 属性，使用标签内的文本
                    if not self.current_a['title'] and self.current_a_text:
                        self.current_a['title'] = self.current_a_text.strip()
                    
                    # 创建 URLItem
                    self.items.append(URLItem(
                        title=self.current_a['title'] or self.current_a['url'],
                        url=self.current_a['url'],
                        category=self.current_a['category'] if self.current_a['category'] else '未分类'
                    ))
                    self.current_a = None
                    self.current_a_text = ""
                    
            def handle_data(self, data):
                if self.in_h3:
                    self.h3_text += data
                elif self.current_a is not None:
                    self.current_a_text += data
        
        content = Path(file_path).read_text(encoding='utf-8')
        parser = BookmarkParser()
        parser.feed(content)
        
        return parser.items


def parse_import_file(file_path: str) -> Tuple[str, List[ImportItem]]:
    """
    根据文件类型自动选择解析器
    
    Args:
        file_path: 文件路径
        
    Returns:
        (文件类型, ImportItem列表)
    """
    path = Path(file_path)
    suffix = path.suffix.lower()
    
    if suffix == '.md':
        return 'markdown', MarkdownParser.parse(file_path)
    elif suffix == '.txt':
        return 'text', TextParser.parse(file_path)
    elif suffix in ['.xlsx', '.xls']:
        return 'excel', ExcelParser.parse(file_path)
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")
